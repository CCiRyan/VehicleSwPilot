import io
import numpy as np
import pytest
from .test_workflow import client, ADMIN, ENGINEER


def upload(client, name, text, **metadata):
    return client.post(
        "/api/knowledge",
        headers=ADMIN,
        files={"file": (name, text.encode())},
        data={"index_mode": "keyword", **metadata},
    )


def test_chapter_boundaries_rule_identifier_and_citation(client):
    text = "# 类型转换\nRule 10.1 使用适当的基本类型。\n规则正文不包含其他章节内容。\n\n# 指针使用\nRule 18.1 指针运算必须保持对象边界。"
    result = upload(client, "coding.md", text, category="misra", version="2012")
    assert result.status_code == 200
    detail = client.get("/api/knowledge/" + result.json()["id"], headers=ADMIN).json()
    assert detail["chunks"] >= 2
    rule = [x for x in detail["items"] if "Rule 10.1" in x["content"]][0]
    assert "18.1" not in rule["content"]
    hits = client.post(
        "/api/knowledge/search",
        headers=ADMIN,
        json={"query": "Rule 10.1", "mode": "keyword"},
    ).json()
    assert "10.1" in hits[0]["content"] and hits[0]["citation_id"].startswith("K-")
    assert hits[0]["category"] == "misra" and hits[0]["version"] == "2012"


def test_duplicate_upload_version_isolation_filters_and_delete(client):
    first = upload(
        client,
        "requirements.md",
        "SWR-001 电压采集采用需求基线。",
        category="srs",
        version="v1",
        asil="B",
    ).json()
    duplicate = upload(
        client,
        "copy.md",
        "SWR-001 电压采集采用需求基线。",
        category="srs",
        version="v1",
        asil="B",
    ).json()
    assert duplicate["duplicate"] is True and duplicate["id"] == first["id"]
    newer = upload(
        client,
        "requirements.md",
        "SWR-002 温度监测采用更新的需求基线。",
        category="srs",
        version="v1",
        asil="B",
    ).json()
    assert newer["revision"] == 2
    hits = client.post(
        "/api/knowledge/search",
        headers=ADMIN,
        json={"query": "SWR-001", "mode": "keyword"},
    ).json()
    assert not hits  # superseded document must not influence latest-context retrieval
    upload(
        client,
        "architecture.md",
        "AUTOSAR 电压采集组件接口。",
        category="autosar",
        version="4.4",
        asil="D",
    )
    filtered = client.post(
        "/api/knowledge/search",
        headers=ADMIN,
        json={
            "query": "电压采集",
            "category": "autosar",
            "asil": "B",
            "mode": "keyword",
        },
    ).json()
    assert filtered == []
    stats = client.get("/api/knowledge/stats", headers=ADMIN).json()
    assert stats["documents"] == 3 and stats["active_documents"] == 2
    assert (
        client.delete("/api/knowledge/" + newer["id"], headers=ADMIN).status_code == 200
    )
    assert client.get("/api/knowledge/" + newer["id"], headers=ADMIN).status_code == 404


def test_keyword_offline_fallback_context_budget_and_embedding_migration(
    client, monkeypatch
):
    import backend.main as main

    response = upload(
        client, "overvoltage.md", "SWR-002 过压检测及输入范围需要验证。", asil="B"
    )
    assert response.status_code == 200
    assert (
        client.post(
            "/api/knowledge/search", headers=ADMIN, json={"query": "过压"}
        ).json()[0]["retrieval"]
        == "keyword"
    )
    assert (
        client.post(
            "/api/knowledge/search",
            headers=ADMIN,
            json={"query": "过压", "mode": "vector"},
        ).status_code
        == 422
    )

    class Embeddings:
        def embed_documents(self, texts):
            return [[1, 0] for _ in texts]

        def embed_query(self, text):
            return [1, 0]

    monkeypatch.setattr(main.knowledge, "embeddings", lambda: Embeddings())
    reindexed = client.post(
        "/api/knowledge/" + response.json()["id"] + "/reindex", headers=ADMIN
    )
    assert reindexed.status_code == 200 and reindexed.json()["index_status"] == "ready"
    result = client.post(
        "/api/knowledge/search", headers=ADMIN, json={"query": "过压", "mode": "hybrid"}
    ).json()
    assert result[0]["retrieval"] == "hybrid" and result[0]["dense_score"] > 0.99
    assert main.knowledge.stats()["cache_entries"] == 1
    monkeypatch.setenv("OPENAI_EMBEDDING_MODEL", "other-model")
    assert main.knowledge.list()[0]["index_status"] == "stale"
    assert (
        client.post(
            "/api/knowledge/search",
            headers=ADMIN,
            json={"query": "过压", "mode": "vector"},
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/knowledge/search", headers=ADMIN, json={"query": "过压"}
        ).json()[0]["retrieval"]
        == "keyword"
    )


def test_invalid_vectors_do_not_mutate_index_and_docx_tables(client, monkeypatch):
    from docx import Document
    import backend.main as main

    class InvalidEmbeddings:
        def embed_documents(self, texts):
            return [[float("nan"), 0] for _ in texts]

    monkeypatch.setattr(main.knowledge, "embeddings", lambda: InvalidEmbeddings())
    response = client.post(
        "/api/knowledge",
        headers=ADMIN,
        files={"file": ("invalid.md", "过压要求".encode())},
    )
    assert response.status_code == 422 and main.knowledge.list() == []
    doc = Document()
    doc.add_heading("接口表", 1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "SWR-003"
    table.cell(0, 1).text = "CAN 超时安全状态"
    stream = io.BytesIO()
    doc.save(stream)
    response = client.post(
        "/api/knowledge",
        headers=ADMIN,
        files={"file": ("interfaces.docx", stream.getvalue())},
        data={"index_mode": "keyword"},
    )
    assert response.status_code == 200
    assert (
        client.post(
            "/api/knowledge/search",
            headers=ADMIN,
            json={"query": "SWR-003", "mode": "keyword"},
        ).json()[0]["source"]
        == "interfaces.docx"
    )
    assert (
        client.post(
            "/api/knowledge",
            headers=ENGINEER,
            files={"file": ("no.md", b"no")},
            data={"index_mode": "keyword"},
        ).status_code
        == 403
    )


def test_dense_filter_cannot_hide_eligible_low_rank_document(client, monkeypatch):
    import backend.main as main

    class Embeddings:
        def embed_documents(self, texts):
            return [[1, 0] if "其他" in text else [0.8, 0.2] for text in texts]

        def embed_query(self, text):
            return [1, 0]

    monkeypatch.setattr(main.knowledge, "embeddings", lambda: Embeddings())
    for i in range(6):
        client.post(
            "/api/knowledge",
            headers=ADMIN,
            files={"file": (f"other{i}.md", f"其他类型文档编号 {i}".encode())},
            data={"category": "other"},
        )
    client.post(
        "/api/knowledge",
        headers=ADMIN,
        files={"file": ("safety.md", "安全状态证据规则".encode())},
        data={"category": "iso26262"},
    )
    hits = client.post(
        "/api/knowledge/search",
        headers=ADMIN,
        json={
            "query": "安全状态",
            "mode": "vector",
            "category": "iso26262",
            "top_k": 1,
        },
    ).json()
    assert hits[0]["source"] == "safety.md"


def test_retrieval_quality_evaluation_reports_misses(client):
    upload(client, "voltage.md", "SWR-002 电压监测的过压检测策略。")
    upload(client, "can.md", "SWR-003 CAN 通信超时处理策略。")
    response = client.post(
        "/api/knowledge/evaluate",
        headers=ADMIN,
        json={
            "mode": "keyword",
            "top_k": 1,
            "cases": [
                {"query": "SWR-002 过压", "expected_sources": ["voltage.md"]},
                {"query": "SWR-003 CAN", "expected_sources": ["can.md"]},
                {"query": "制动压力冗余", "expected_sources": ["unavailable.md"]},
            ],
        },
    )
    assert response.status_code == 200
    result = response.json()
    assert result["hit_rate"] == pytest.approx(2 / 3)
    assert result["mrr"] == pytest.approx(2 / 3)
    assert result["cases"][2]["recall_at_k"] == 0
