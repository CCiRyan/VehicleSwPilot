"""Versioned, section-aware RAG with Faiss + BM25, source filters and RRF/MMR."""

import hashlib
import io
import json
import math
import re
import threading
import unicodedata
import uuid
from collections import Counter, OrderedDict
import numpy as np
from .models import embedding_model, embedding_signature
from .store import now


def tokens(text):
    text = unicodedata.normalize("NFKC", text).lower()
    result = re.findall(r"[a-z0-9]+(?:[_.:/-][a-z0-9]+)*", text)
    result += ["rule" + m for m in re.findall(r"(?:rule|规则)\s*(\d+\.\d+)", text)]
    for segment in re.findall(r"[\u4e00-\u9fff]+", text):
        result += (
            [segment[i : i + 2] for i in range(len(segment) - 1)]
            if len(segment) > 1
            else [segment]
        )
    return result


def clean_text(text):
    text = unicodedata.normalize("NFKC", text).replace("\x00", "")
    return "\n".join(
        re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()
    ).strip()


def parse_pages(name, body):
    suffix = name.lower().rsplit(".", 1)[-1]
    if suffix == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(body))
        if reader.is_encrypted:
            raise ValueError("加密 PDF 请解密后上传")
        if len(reader.pages) > 2000:
            raise ValueError("PDF 超过 2000 页，请按章节拆分")
        return [
            (i + 1, clean_text(p.extract_text() or ""))
            for i, p in enumerate(reader.pages)
        ]
    if suffix == "docx":
        from docx import Document

        document = Document(io.BytesIO(body))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            paragraphs += [
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            ]
        return [(1, clean_text("\n".join(paragraphs)))]
    if suffix not in ("txt", "md"):
        raise ValueError("仅支持 PDF、DOCX、TXT、Markdown")
    try:
        return [(1, clean_text(body.decode("utf-8-sig")))]
    except UnicodeDecodeError as exc:
        raise ValueError("文本编码必须为 UTF-8") from exc


def section_chunks(pages, size=900, overlap=100):
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", "。", "；", ". ", " ", ""],
    )
    heading = "文档正文"
    chunks = []
    pattern = re.compile(
        r"^(?:#{1,6}\s+.+|\d+(?:\.\d+){0,5}[ \t]+.{1,100}|(?:Rule|Directive|规则|条款)\s*\d+(?:\.\d+)+.*|(?:SYS|SWR|SRS|DES|TC)-[A-Za-z0-9_-]+.*)$",
        re.I,
    )
    for page, text in pages:
        sections = []
        lines = []
        for line in text.splitlines():
            if pattern.match(line):
                if lines:
                    sections.append((heading, "\n".join(lines)))
                    lines = []
                heading = line.lstrip("# ").strip()
            lines.append(line)
        if lines:
            sections.append((heading, "\n".join(lines)))
        for section, body in sections:
            for content in splitter.split_text(body):
                if not content.strip():
                    continue
                identifiers = re.findall(
                    r"\b(?:SYS|SWR|SRS|ARC|DES|SRC|UT|IT|ST|AT|TC)-[A-Za-z0-9_-]+",
                    content,
                )
                rules = re.findall(
                    r"(?:Rule|Directive|规则|条款)\s*\d+(?:\.\d+)+", content, re.I
                )
                chunks.append(
                    dict(
                        page=page,
                        section=section,
                        content=content,
                        identifiers=sorted(set(identifiers + rules)),
                    )
                )
    return chunks


class Knowledge:
    def __init__(self, store):
        self.store = store
        self.lock = threading.RLock()
        self.cache = OrderedDict()
        self.reload()

    def embeddings(self):
        return embedding_model()

    def reload(self):
        import faiss

        with self.lock, self.store.lock:
            docs = [
                json.loads(row[0])
                for row in self.store.conn.execute(
                    "SELECT data FROM knowledge ORDER BY rowid"
                )
            ]
            chunks = []
            groups = {}
            for doc in docs:
                if not doc.get("active", True):
                    continue
                for original in doc["items"]:
                    item = {
                        **original,
                        "document_id": doc["id"],
                        "source": doc["name"],
                        **doc.get("metadata", {}),
                    }
                    item.setdefault("section", "文档正文")
                    item.setdefault("chunk_id", f'{doc["id"]}:{item["chunk"]}')
                    item.setdefault("citation_id", "K-" + item["chunk_id"])
                    item.setdefault("identifiers", [])
                    item["embedding_signature"] = doc.get(
                        "embedding_signature", "legacy|unknown"
                    )
                    chunks.append(item)
                    if item.get("vector"):
                        groups.setdefault(
                            (item["embedding_signature"], len(item["vector"])), []
                        ).append(len(chunks) - 1)
            indexes = {}
            for signature, indices in groups.items():
                matrix = np.asarray(
                    [chunks[i]["vector"] for i in indices], dtype="float32"
                )
                if not np.isfinite(matrix).all():
                    continue
                faiss.normalize_L2(matrix)
                index = faiss.IndexFlatIP(matrix.shape[1])
                index.add(matrix)
                indexes[signature] = (index, indices)
            self.docs = docs
            self.chunks = chunks
            self.indexes = indexes
            self.term_counts = [
                Counter(tokens(c["content"] + " " + c["section"])) for c in chunks
            ]
            self.lengths = [sum(c.values()) for c in self.term_counts]
            self.df = Counter(term for counts in self.term_counts for term in counts)
            self.average_length = sum(self.lengths) / max(len(chunks), 1)
            self.cache.clear()
            current = [
                index
                for (signature, _), (index, _) in indexes.items()
                if signature == embedding_signature()
            ]
            self.index = current[0] if current else None

    def _validate_vectors(self, vectors, count):
        array = np.asarray(vectors, dtype="float32")
        if (
            array.ndim != 2
            or array.shape[0] != count
            or array.shape[1] < 1
            or not np.isfinite(array).all()
            or (np.linalg.norm(array, axis=1) == 0).any()
        ):
            raise ValueError("Embedding 返回数量、维度或数值无效；未写入知识库")
        return array.tolist()

    def add(
        self,
        name,
        body,
        metadata=None,
        index_mode="vector",
        chunk_size=900,
        chunk_overlap=100,
    ):
        if index_mode not in ("vector", "keyword"):
            raise ValueError("索引模式必须为 vector 或 keyword")
        if not 200 <= chunk_size <= 2000 or not 0 <= chunk_overlap < min(
            chunk_size, 500
        ):
            raise ValueError("切片大小需为 200–2000，重叠需小于切片大小且小于 500")
        metadata = metadata or {}
        category = metadata.get("category", "other")
        if category not in ("autosar", "misra", "iso26262", "srs", "project", "other"):
            raise ValueError("未知知识类别")
        metadata = {
            key: str(metadata.get(key, ""))[:200]
            for key in ["category", "version", "asil", "stage"]
        }
        metadata["category"] = category
        fingerprint = hashlib.sha256(body).hexdigest()
        with self.lock:
            for doc in self.docs:
                if (
                    doc.get("content_hash") == fingerprint
                    and doc.get("metadata", {}) == metadata
                    and doc.get("active", True)
                    and doc.get("index_mode", "vector") == index_mode
                ):
                    return {**self._public(doc), "duplicate": True}
        items = section_chunks(parse_pages(name, body), chunk_size, chunk_overlap)
        if not items:
            raise ValueError("文档没有可提取文本；扫描 PDF 请先进行 OCR。")
        if len(items) > 1000:
            raise ValueError("切片数量超过 1000，请拆分文档。")
        if index_mode == "vector":
            vectors = []
            embedder = self.embeddings()
            for offset in range(0, len(items), 64):
                texts = [
                    c["section"] + "\n" + c["content"]
                    for c in items[offset : offset + 64]
                ]
                vectors += self._validate_vectors(
                    embedder.embed_documents(texts), len(texts)
                )
            if len({len(v) for v in vectors}) != 1:
                raise ValueError("Embedding 批次返回维度不一致")
        else:
            vectors = [None] * len(items)
        doc_id = uuid.uuid4().hex
        for i, (item, vector) in enumerate(zip(items, vectors)):
            item.update(
                source=name,
                chunk=i + 1,
                chunk_id=f"{doc_id}:{i+1}",
                citation_id=f"K-{doc_id}:{i+1}",
                vector=vector,
            )
        doc = dict(
            id=doc_id,
            name=name,
            chunks=len(items),
            created_at=now(),
            items=items,
            metadata=metadata,
            content_hash=fingerprint,
            embedding_signature=(
                embedding_signature() if index_mode == "vector" else None
            ),
            index_mode=index_mode,
            active=True,
            revision=1,
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
        )
        with self.lock, self.store.lock, self.store.conn:
            latest = [
                json.loads(row[0])
                for row in self.store.conn.execute("SELECT data FROM knowledge")
            ]
            for old in latest:
                if (
                    old.get("content_hash") == fingerprint
                    and old.get("metadata", {}) == metadata
                    and old.get("active", True)
                    and old.get("index_mode", "vector") == index_mode
                ):
                    return {**self._public(old), "duplicate": True}
            for old in latest:
                if (
                    old["name"] == name
                    and old.get("metadata", {}) == metadata
                    and old.get("active", True)
                ):
                    doc["revision"] = max(doc["revision"], old.get("revision", 1) + 1)
                    old["active"] = False
                    self.store.conn.execute(
                        "UPDATE knowledge SET data=? WHERE id=?",
                        (json.dumps(old, ensure_ascii=False), old["id"]),
                    )
            self.store.conn.execute(
                "INSERT INTO knowledge VALUES(?,?)",
                (doc_id, json.dumps(doc, ensure_ascii=False)),
            )
        self.reload()
        return self._public(doc)

    def _public(self, doc):
        result = {k: v for k, v in doc.items() if k != "items"}
        result["metadata"] = doc.get("metadata", {})
        result.setdefault("active", True)
        result["index_status"] = (
            "keyword"
            if doc.get("index_mode") == "keyword"
            else (
                "ready"
                if doc.get("embedding_signature") == embedding_signature()
                else "stale"
            )
        )
        return result

    def list(self):
        with self.lock:
            return [self._public(doc) for doc in self.docs]

    def delete(self, document_id):
        with self.lock, self.store.lock, self.store.conn:
            if not self.store.conn.execute(
                "DELETE FROM knowledge WHERE id=?", (document_id,)
            ).rowcount:
                raise KeyError(document_id)
        self.reload()

    def detail(self, document_id):
        with self.lock:
            doc = next((d for d in self.docs if d["id"] == document_id), None)
            if doc is None:
                raise KeyError(document_id)
            return {
                **self._public(doc),
                "items": [
                    {k: v for k, v in item.items() if k != "vector"}
                    for item in doc["items"]
                ],
            }

    def reindex(self, document_id):
        with self.lock:
            doc = next((d for d in self.docs if d["id"] == document_id), None)
            if doc is None:
                raise KeyError(document_id)
            snapshot = json.loads(json.dumps(doc))
        embedder = self.embeddings()
        items = snapshot["items"]
        vectors = []
        for offset in range(0, len(items), 64):
            texts = [
                c.get("section", "") + "\n" + c["content"]
                for c in items[offset : offset + 64]
            ]
            vectors += self._validate_vectors(
                embedder.embed_documents(texts), len(texts)
            )
        if len({len(v) for v in vectors}) != 1:
            raise ValueError("Embedding 批次返回维度不一致")
        for item, vector in zip(items, vectors):
            item["vector"] = vector
        snapshot.update(
            embedding_signature=embedding_signature(),
            index_mode="vector",
            reindexed_at=now(),
        )
        with self.lock, self.store.lock, self.store.conn:
            current = self.store.conn.execute(
                "SELECT data FROM knowledge WHERE id=?", (document_id,)
            ).fetchone()
            if not current:
                raise KeyError(document_id)
            snapshot["active"] = json.loads(current[0]).get("active", True)
            self.store.conn.execute(
                "UPDATE knowledge SET data=? WHERE id=?",
                (json.dumps(snapshot, ensure_ascii=False), document_id),
            )
        self.reload()
        return self._public(snapshot)

    def stats(self):
        with self.lock:
            return dict(
                documents=len(self.docs),
                active_documents=sum(d.get("active", True) for d in self.docs),
                chunks=len(self.chunks),
                vector_groups=len(self.indexes),
                stale_documents=sum(
                    self._public(d)["index_status"] == "stale" and d.get("active", True)
                    for d in self.docs
                ),
                embedding_model=embedding_signature().split("|")[1],
                cache_entries=len(self.cache),
                retrieval="BM25 + Faiss / RRF + diversity",
            )

    def search(
        self,
        query,
        top_k=4,
        filters=None,
        mode="hybrid",
        min_score=0.25,
        context_chars=10000,
    ):
        if not query.strip():
            return []
        filters = filters or {}
        if mode not in ("hybrid", "keyword", "vector"):
            raise ValueError("未知检索模式")
        warnings = []
        with self.lock:
            chunks = list(self.chunks)
            indexes = dict(self.indexes)
            counts = self.term_counts
            lengths = self.lengths
            df = self.df
            average = self.average_length
            eligible = [
                i
                for i, c in enumerate(chunks)
                if all(
                    not value
                    or (
                        c.get(key, "") in (value, "", "all")
                        if key in ("asil", "stage")
                        else c.get(key) == value
                    )
                    for key, value in filters.items()
                )
            ]
            if not eligible:
                return []
            current = [
                (signature, data)
                for signature, data in indexes.items()
                if signature[0] == embedding_signature()
            ]
        lexical = {}
        terms = Counter(tokens(query))
        N = len(chunks)
        for i in eligible:
            score = 0.0
            for term in terms:
                tf = counts[i].get(term, 0)
                if tf:
                    idf = math.log(1 + (N - df[term] + 0.5) / (df[term] + 0.5))
                    score += (
                        idf
                        * tf
                        * 2.2
                        / (tf + 1.2 * (0.25 + 0.75 * lengths[i] / max(average, 1)))
                    )
            if score > 0:
                lexical[i] = score
        dense = {}
        if mode != "keyword" and current:
            try:
                cache_key = (embedding_signature(), query)
                with self.lock:
                    vector = self.cache.get(cache_key)
                if vector is None:
                    vector = self.embeddings().embed_query(query)
                    with self.lock:
                        self.cache[cache_key] = vector
                        while len(self.cache) > 128:
                            self.cache.popitem(last=False)
                q = self._validate_vectors([vector], 1)
                import faiss

                normalized = np.asarray(q, dtype="float32")
                faiss.normalize_L2(normalized)
                for (signature, dimension), (index, indices) in current:
                    if dimension != normalized.shape[1]:
                        warnings.append("Embedding 维度发生变化，需要重建索引")
                        continue
                    index_ids = set(indices)
                    subset = [i for i in eligible if i in index_ids]
                    if not subset:
                        continue
                    if len(subset) == len(indices):
                        filtered = index
                        subset = indices
                    else:
                        matrix = np.asarray(
                            [chunks[i]["vector"] for i in subset], dtype="float32"
                        )
                        faiss.normalize_L2(matrix)
                        filtered = faiss.IndexFlatIP(dimension)
                        filtered.add(matrix)
                    scores, ids = filtered.search(
                        normalized, min(len(subset), max(top_k * 8, 32))
                    )
                    for score, local in zip(scores[0], ids[0]):
                        if local >= 0 and score >= min_score:
                            dense[subset[int(local)]] = float(score)
            except Exception:
                if mode == "vector":
                    raise
                warnings.append("向量检索不可用，已降级为关键词检索")
        elif mode == "vector":
            raise ValueError("没有匹配当前 Embedding 配置的向量索引，请导入或重建索引")
        elif mode == "hybrid":
            warnings.append("当前未启用匹配的向量索引，使用关键词检索")
        if mode == "vector":
            lexical = {}
        fused = {}
        for ranking in [
            sorted(lexical, key=lexical.get, reverse=True),
            sorted(dense, key=dense.get, reverse=True),
        ]:
            for rank, i in enumerate(ranking):
                fused[i] = fused.get(i, 0) + 1 / (60 + rank + 1)
        exact = set(
            re.findall(
                r"\b(?:SYS|SWR|SRS|ARC|DES|SRC|UT|IT|ST|AT|TC)-[A-Za-z0-9_-]+",
                query,
                re.I,
            )
        )
        exact.update(
            re.findall(r"(?:Rule|Directive|规则|条款)\s*\d+(?:\.\d+)+", query, re.I)
        )
        for i in fused:
            if any(
                identifier.lower() in chunks[i]["content"].lower()
                for identifier in exact
            ):
                fused[i] += 0.04
        selected = []
        remaining = set(fused)
        budget = 0
        fingerprints = set()
        results = []
        termsets = {i: set(tokens(chunks[i]["content"])) for i in fused}
        while remaining and len(results) < top_k:

            def priority(i):
                similarity = max(
                    (
                        len(termsets[i] & termsets[j])
                        / max(len(termsets[i] | termsets[j]), 1)
                        for j in selected
                    ),
                    default=0,
                )
                same_doc = sum(
                    chunks[j]["document_id"] == chunks[i]["document_id"]
                    for j in selected
                )
                return fused[i] * (1 - 0.45 * similarity) / (1 + 0.12 * same_doc)

            i = max(remaining, key=lambda n: (priority(n), -n))
            remaining.remove(i)
            chunk = chunks[i]
            fingerprint = hashlib.sha256(chunk["content"].encode()).hexdigest()
            if (
                fingerprint in fingerprints
                or budget + len(chunk["content"]) > context_chars
            ):
                continue
            fingerprints.add(fingerprint)
            selected.append(i)
            budget += len(chunk["content"])
            results.append(
                {
                    **{
                        k: v
                        for k, v in chunk.items()
                        if k not in ("vector", "embedding_signature")
                    },
                    "score": dense.get(i, lexical.get(i, 0) / (1 + lexical.get(i, 0))),
                    "dense_score": dense.get(i),
                    "keyword_score": lexical.get(i, 0),
                    "relevance": min(1.0, fused[i] * 30.5),
                    "retrieval": (
                        "hybrid"
                        if i in dense and i in lexical
                        else "vector" if i in dense else "keyword"
                    ),
                    "warnings": warnings.copy(),
                }
            )
        for category in ("misra", "autosar", "iso26262"):
            versions = {
                r.get("version")
                for r in results
                if r.get("category") == category and r.get("version")
            }
            if len(versions) > 1:
                for result in results:
                    result["warnings"].append(
                        f"{category} 检索命中多个规范版本，请设置版本过滤并核对适用性"
                    )
        return results
